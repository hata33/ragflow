#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import logging
from io import BytesIO
from pptx import Presentation


class RAGFlowPptParser:
    """PowerPoint 解析器，按页面顺序提取文本、表格和分组形状内容。"""

    def __init__(self):
        super().__init__()
        self._shape_cache = {}

    def __sort_shapes(self, shapes):
        """按近似阅读顺序排序形状，避免抽取结果在视觉上跳跃。"""

        cache_key = id(shapes)
        if cache_key not in self._shape_cache:
            self._shape_cache[cache_key] = sorted(
                shapes, 
                key=lambda x: ((x.top if x.top is not None else 0) // 10, x.left if x.left is not None else 0)
            )
        return self._shape_cache[cache_key]

    def __get_bulleted_text(self, paragraph):
        """保留项目符号层级信息，便于还原列表结构。"""

        is_bulleted = bool(paragraph._p.xpath("./a:pPr/a:buChar")) or bool(paragraph._p.xpath("./a:pPr/a:buAutoNum")) or bool(paragraph._p.xpath("./a:pPr/a:buBlip"))
        if is_bulleted:
            return f"{'  '* paragraph.level}.{paragraph.text}"
        else:
            return paragraph.text

    def __extract(self, shape):
        """递归提取单个形状中的文本内容。"""

        try:
            # 优先走文本框分支，这是最常见也最稳定的文本来源。
            if hasattr(shape, 'has_text_frame') and shape.has_text_frame:
                text_frame = shape.text_frame
                texts = []
                for paragraph in text_frame.paragraphs:
                    if paragraph.text.strip():
                        texts.append(self.__get_bulleted_text(paragraph))
                return "\n".join(texts)

            # 某些形状访问 `shape_type` 时会抛异常，因此这里单独兜底。
            try:
                shape_type = shape.shape_type
            except NotImplementedError:
                # 如果拿不到类型，尽量退化为直接读取 shape.text。
                if hasattr(shape, 'text'):
                    return shape.text.strip()
                return ""

            # 表格按“表头: 单元格内容”的形式转成结构化文本。
            if shape_type == 19:
                tb = shape.table
                rows = []
                for i in range(1, len(tb.rows)):
                    rows.append("; ".join([tb.cell(
                        0, j).text + ": " + tb.cell(i, j).text for j in range(len(tb.columns)) if tb.cell(i, j)]))
                return "\n".join(rows)

            # 分组形状需要递归展开，否则会漏掉内部文本。
            if shape_type == 6:
                texts = []
                for p in self.__sort_shapes(shape.shapes):
                    t = self.__extract(p)
                    if t:
                        texts.append(t)
                return "\n".join(texts)

            return ""

        except Exception as e:
            logging.error(f"Error processing shape: {str(e)}")
            return ""

    def __call__(self, fnm, from_page, to_page, callback=None):
        """按页提取 PPT 内容，返回每页拼接后的文本列表。"""

        ppt = Presentation(fnm) if isinstance(
            fnm, str) else Presentation(
            BytesIO(fnm))
        txts = []
        self.total_page = len(ppt.slides)
        for i, slide in enumerate(ppt.slides):
            if i < from_page:
                continue
            if i >= to_page:
                break
            texts = []
            for shape in self.__sort_shapes(slide.shapes):
                txt = self.__extract(shape)
                if txt:
                    texts.append(txt)
            txts.append("\n".join(texts))

        return txts
