#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

from docx import Document
import re
import pandas as pd
from collections import Counter
from rag.nlp import rag_tokenizer
from io import BytesIO
import logging
from common.constants import MAXIMUM_PAGE_NUMBER
from docx.image.exceptions import (
    InvalidImageStreamError,
    UnexpectedEndOfFileError,
    UnrecognizedImageError,
)
from rag.utils.lazy_image import LazyImage


class RAGFlowDocxParser:
    """Word 解析器，负责提取段落、表格以及内嵌图片。"""

    def get_picture(self, document, paragraph):
        """提取段落中的嵌入图片，并延迟包装成 `LazyImage`。"""

        imgs = paragraph._element.xpath(".//pic:pic")
        if not imgs:
            return None
        image_blobs = []
        for img in imgs:
            embed = img.xpath(".//a:blip/@r:embed")
            if not embed:
                continue
            embed = embed[0]
            image_blob = None
            try:
                related_part = document.part.related_parts[embed]
            except Exception as e:
                logging.warning(f"Skipping image due to unexpected error getting related_part: {e}")
                continue

            try:
                # 优先通过 python-docx 暴露的标准 image 接口取图。
                image = related_part.image
                if image is not None:
                    image_blob = image.blob
            except (
                UnrecognizedImageError,
                UnexpectedEndOfFileError,
                InvalidImageStreamError,
                UnicodeDecodeError,
            ) as e:
                logging.info(f"Damaged image encountered, attempting blob fallback: {e}")
            except Exception as e:
                logging.warning(f"Unexpected error getting image, attempting blob fallback: {e}")

            if image_blob is None:
                # 某些损坏文档下标准接口会失败，这里退化为直接拿底层 blob。
                image_blob = getattr(related_part, "blob", None)
            if image_blob:
                image_blobs.append(image_blob)
        if not image_blobs:
            return None
        return LazyImage(image_blobs)


    def __extract_table_content(self, tb):
        """将 Word 表格转成 DataFrame 后复用统一表格整理逻辑。"""

        df = []
        for row in tb.rows:
            df.append([c.text for c in row.cells])
        return self.__compose_table_content(pd.DataFrame(df))

    def __compose_table_content(self, df):
        """根据单元格内容类型推断表头，并把表格转成可检索文本。"""

        def blockType(b):
            # 通过规则快速判断单元格更像日期、数值、编号还是长文本。
            pattern = [
                ("^(20|19)[0-9]{2}[年/-][0-9]{1,2}[月/-][0-9]{1,2}日*$", "Dt"),
                (r"^(20|19)[0-9]{2}年$", "Dt"),
                (r"^(20|19)[0-9]{2}[年/-][0-9]{1,2}月*$", "Dt"),
                ("^[0-9]{1,2}[月/-][0-9]{1,2}日*$", "Dt"),
                (r"^第*[一二三四1-4]季度$", "Dt"),
                (r"^(20|19)[0-9]{2}年*[一二三四1-4]季度$", "Dt"),
                (r"^(20|19)[0-9]{2}[ABCDE]$", "DT"),
                ("^[0-9.,+%/ -]+$", "Nu"),
                (r"^[0-9A-Z/\._~-]+$", "Ca"),
                (r"^[A-Z]*[a-z' -]+$", "En"),
                (r"^[0-9.,+-]+[0-9A-Za-z/$￥%<>（）()' -]+$", "NE"),
                (r"^.{1}$", "Sg")
            ]
            for p, n in pattern:
                if re.search(p, b):
                    return n
            tks = [t for t in rag_tokenizer.tokenize(b).split() if len(t) > 1]
            if len(tks) > 3:
                if len(tks) < 12:
                    return "Tx"
                else:
                    return "Lx"

            if len(tks) == 1 and rag_tokenizer.tag(tks[0]) == "nr":
                return "Nr"

            return "Ot"

        if len(df) < 2:
            return []
        max_type = Counter([blockType(str(df.iloc[i, j])) for i in range(
            1, len(df)) for j in range(len(df.iloc[i, :]))])
        max_type = max(max_type.items(), key=lambda x: x[1])[0]

        colnm = len(df.iloc[0, :])
        # 表头不一定只占第一行，这里允许继续向下吸收“表头风格”的行。
        hdrows = [0]
        if max_type == "Nu":
            for r in range(1, len(df)):
                tys = Counter([blockType(str(df.iloc[r, j]))
                              for j in range(len(df.iloc[r, :]))])
                tys = max(tys.items(), key=lambda x: x[1])[0]
                if tys != max_type:
                    hdrows.append(r)

        lines = []
        for i in range(1, len(df)):
            if i in hdrows:
                continue
            hr = [r - i for r in hdrows]
            hr = [r for r in hr if r < 0]
            t = len(hr) - 1
            while t > 0:
                if hr[t] - hr[t - 1] > 1:
                    hr = hr[t:]
                    break
                t -= 1
            headers = []
            for j in range(len(df.iloc[i, :])):
                t = []
                for h in hr:
                    x = str(df.iloc[i + h, j]).strip()
                    if x in t:
                        continue
                    t.append(x)
                t = ",".join(t)
                if t:
                    t += ": "
                headers.append(t)
            cells = []
            for j in range(len(df.iloc[i, :])):
                if not str(df.iloc[i, j]):
                    continue
                cells.append(headers[j] + str(df.iloc[i, j]))
            lines.append(";".join(cells))

        if colnm > 3:
            return lines
        return ["\n".join(lines)]

    def __call__(self, fnm, from_page=0, to_page=MAXIMUM_PAGE_NUMBER):
        self.doc = Document(fnm) if isinstance(
            fnm, str) else Document(BytesIO(fnm))
        pn = 0  # 当前已解析到的近似页码。
        secs = []  # 收集段落文本与样式名。
        for p in self.doc.paragraphs:
            if pn > to_page:
                break

            runs_within_single_paragraph = []  # 只保留当前页范围内的 run 内容。
            for run in p.runs:
                if pn > to_page:
                    break
                if from_page <= pn < to_page and p.text.strip():
                    runs_within_single_paragraph.append(run.text)

                # Word 文档没有稳定页码概念，这里只能通过渲染分页符近似累计。
                if 'lastRenderedPageBreak' in run._element.xml:
                    pn += 1

            secs.append(("".join(runs_within_single_paragraph), p.style.name if hasattr(p.style, 'name') else ''))

        tbls = [self.__extract_table_content(tb) for tb in self.doc.tables]
        return secs, tbls
